"""Word documents: research reports and general documents written by the LLM."""
import datetime
import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

from .. import config, ui
from ..llm import router
from .registry import tool

ACCENT = RGBColor(0x0B, 0x5C, 0xAD)


def slug(s, n=60):
    return re.sub(r"[^a-zA-Z0-9]+", "_", s).strip("_")[:n] or "document"


def _base_doc():
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1)
    for lvl, size in ((1, 16), (2, 13), (3, 12)):
        h = d.styles[f"Heading {lvl}"]
        h.font.color.rgb, h.font.size, h.font.name = ACCENT, Pt(size), "Calibri"
    return d


def add_hyperlink(par, url, text):
    part = par.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0B5CAD"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); run.append(t)
    link.append(run)
    par._p.append(link)


def add_rich(par, text):
    """Supports **bold** and *italic* inline markdown."""
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if chunk.startswith("**") and chunk.endswith("**"):
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            par.add_run(chunk[1:-1]).italic = True
        else:
            par.add_run(chunk)


def add_markdown(d, md):
    for block in md.split("\n"):
        line = block.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            d.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif re.match(r"^\s*[-*•]\s+", line):
            add_rich(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            add_rich(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        else:
            add_rich(d.add_paragraph(), line)


def _footer(d):
    p = d.sections[0].footer.paragraphs[0]
    p.text = f"Prepared by JARVIS for {config.USER_NAME} · {datetime.date.today():%B %d, %Y}"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(8)


def research_report(topic, papers, syn, stats, style, min_cred):
    d = _base_doc()
    t = d.add_heading(syn.get("title") or f"Research Brief: {topic}", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = d.add_paragraph(f"{datetime.date.today():%B %d, %Y} · {len(papers)} peer-reviewed sources · "
                          f"credibility threshold {min_cred}/100 · {style} citations")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _footer(d)

    d.add_heading("Executive summary", 1)
    add_markdown(d, syn.get("executive_summary", ""))

    for s in syn.get("sections", []):
        d.add_heading(s["heading"], 1)
        for para in s["body"].split("\n\n"):
            add_markdown(d, para)

    for key, head in (("consensus", "Where the evidence agrees"), ("disagreements", "Where it conflicts"),
                      ("gaps", "Gaps and weaknesses in the evidence"), ("implications", "Practical implications"),
                      ("conclusion", "Conclusion")):
        if syn.get(key):
            d.add_heading(head, 1)
            add_markdown(d, syn[key])

    d.add_page_break()
    d.add_heading("Source summaries", 1)
    for i, p in enumerate(papers, 1):
        d.add_heading(f"[{i}] {p['title']}", 2)
        meta = d.add_paragraph()
        meta.add_run(f"{p.get('authors', '')} ({p.get('year')}). {p.get('venue', '')}. ").italic = True
        meta.add_run(f"Credibility {p['credibility']}/100 · {p.get('citations', 0)} citations · "
                     f"summarised from {p.get('summary_basis', 'abstract')} · via {p.get('source')}")
        link = f"https://doi.org/{p['doi']}" if p.get("doi") else p.get("url")
        if link:
            add_hyperlink(d.add_paragraph(), link, link)
        add_markdown(d, p.get("summary", ""))
        if p.get("key_findings"):
            d.add_paragraph().add_run("Key findings").bold = True
            for f in p["key_findings"]:
                d.add_paragraph(str(f), style="List Bullet")
        for label, key in (("Methods", "methods"), ("Limitations", "limitations"), ("Relevance", "relevance_note")):
            if p.get(key):
                par = d.add_paragraph()
                par.add_run(f"{label}: ").bold = True
                par.add_run(str(p[key]))

    d.add_heading("Credibility scoring", 1)
    d.add_paragraph("Each source is scored 0–100: citation impact (35, adjusted for age), venue quality (25), "
                    "verification signals such as DOI and full text (15), recency (15), and depth of available content (10). "
                    f"Retracted work scores 0. Only sources scoring {min_cred}+ that an AI relevance check rated 6/10 or higher were kept. "
                    f"{stats.get('examined', 0)} candidate papers were examined using the queries: "
                    + "; ".join(stats.get("queries", [])) + ".")
    tbl = d.add_table(rows=1, cols=6)
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(tbl.rows[0].cells, ["#", "Source", "Score", "Citations", "Venue", "Recency"]):
        c.text = h
    for i, p in enumerate(papers, 1):
        b = p.get("credibility_breakdown", {})
        row = tbl.add_row().cells
        vals = [str(i), f"{p['title'][:60]}", str(p["credibility"]), f"{b.get('citations', '')}/35",
                f"{b.get('venue', '')}/25", f"{b.get('recency', '')}/15"]
        for c, v in zip(row, vals):
            c.text = v

    d.add_heading("References", 1)
    for i, p in enumerate(papers, 1):
        par = d.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.4)
        par.paragraph_format.first_line_indent = Inches(-0.4)
        par.add_run(f"[{i}] {p['citation']}")

    out = config.OUTPUT_DIR / f"Research_{slug(topic)}_{datetime.datetime.now():%Y%m%d_%H%M}.docx"
    d.save(out)
    return str(out)


@tool("Write a document (essay, report, notes, letter, study guide...) and save it as a Word file. "
      "Optionally base it on a file or research report.",
      title="document title", instructions="what the document should contain, audience, length",
      source_file="optional path of a file to base it on")
def write_document(title, instructions, source_file=""):
    from .files import extract_text
    src = ""
    if source_file:
        src = "\n\nSOURCE MATERIAL:\n" + extract_text(source_file)[:150000]
    ui.status(f"Drafting '{title}'")
    body = router().ask(
        f"Write a complete, well-structured document titled '{title}'.\nInstructions: {instructions}{src}\n\n"
        "Use markdown: ## headings, bullet lists, **bold** for key terms. If you use source material, cite it. "
        "Do not include the title line itself.",
        system="You are an excellent writer and careful thinker. Be accurate, specific and well organised.",
        task="document", temperature=0.5)
    d = _base_doc()
    d.add_heading(title, 0)
    _footer(d)
    add_markdown(d, body)
    out = config.OUTPUT_DIR / f"{slug(title)}_{datetime.datetime.now():%Y%m%d_%H%M}.docx"
    d.save(out)
    return f"Document saved: {out}"
