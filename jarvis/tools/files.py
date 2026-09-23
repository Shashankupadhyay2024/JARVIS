"""Reading and analysing any file: PDF, Word, PowerPoint, Excel/CSV, images, audio/video, code, text."""
from pathlib import Path

from .. import config
from ..llm import router
from ..safety import confirmer
from .registry import tool

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic", ".bmp", ".tiff"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".mp4", ".mov", ".mkv", ".webm"}


def resolve(path):
    p = Path(str(path).strip().strip("'\"")).expanduser()
    if not p.is_absolute():
        for base in (config.INBOX_DIR, Path.home() / "Desktop", Path.home() / "Downloads", Path.home() / "Documents", Path.cwd()):
            if (base / p).exists():
                return base / p
    return p


def extract_text(path, max_chars=400_000):
    p = resolve(path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {p}")
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(str(p))
        text = "\n\n".join(f"[page {i + 1}]\n{pg.extract_text() or ''}" for i, pg in enumerate(r.pages))
    elif ext == ".docx":
        import docx
        d = docx.Document(str(p))
        text = "\n".join(par.text for par in d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                text += "\n" + " | ".join(c.text for c in row.cells)
    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(p))
        out = []
        for i, s in enumerate(prs.slides, 1):
            texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
            notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
            out.append(f"[slide {i}]\n" + "\n".join(texts) + (f"\nNotes: {notes}" if notes else ""))
        text = "\n\n".join(out)
    elif ext in (".xlsx", ".xls", ".csv", ".tsv"):
        import pandas as pd
        if ext in (".csv", ".tsv"):
            sheets = {"data": pd.read_csv(p, sep="\t" if ext == ".tsv" else ",")}
        else:
            sheets = pd.read_excel(p, sheet_name=None)
        out = []
        for name, df in sheets.items():
            out.append(f"[sheet {name}] {df.shape[0]} rows x {df.shape[1]} cols\nColumns: {list(df.columns)}\n"
                       f"Summary:\n{df.describe(include='all').to_string()[:4000]}\nFirst rows:\n{df.head(25).to_string()}")
        text = "\n\n".join(out)
    elif ext in AUDIO_EXT:
        from ..voice import transcribe
        text = "[transcript]\n" + transcribe(str(p))
    elif ext in IMAGE_EXT:
        text = "[image file]"
    else:
        text = p.read_text(errors="ignore")
    return text[:max_chars]


@tool("List files in a folder (defaults to the JARVIS inbox where you can drop files).", path="folder path")
def list_dir(path=str(config.INBOX_DIR)):
    p = resolve(path)
    if not p.is_dir():
        return f"Not a folder: {p}"
    items = sorted(p.iterdir(), key=lambda x: x.stat().st_mtime, reverse=True)[:60]
    return "\n".join(f"{'📁' if i.is_dir() else '📄'} {i}" for i in items) or "(empty)"


@tool("Read the raw text of a file (first 15k chars).", path="file path")
def read_file(path):
    t = extract_text(path)
    return t[:15000] + ("\n...[truncated]" if len(t) > 15000 else "")


@tool("Write text to a file. Asks before overwriting an existing file.", path="file path", content="text")
def write_file(path, content):
    p = resolve(path)
    if p.exists() and not confirmer.confirm(f"Overwrite {p}?"):
        return "User declined."
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content)
    return f"Wrote {len(content)} chars to {p}"


def _chunks(text, size):
    return [text[i:i + size] for i in range(0, len(text), size)]


@tool("Analyse any file (PDF, Word, PowerPoint, Excel/CSV, image, audio/video, code) and answer a question or summarise it.",
      path="file path (drag the file into Terminal to paste its path)", question="what you want to know")
def analyze_file(path, question="Summarise this file: purpose, key points, notable details, and anything that looks wrong."):
    p = resolve(path)
    llm = router()
    system = ("You are an expert analyst. Answer precisely using only the file's content, quote page/slide/row "
              "references when useful, and point out errors, weak arguments or data problems you notice.")
    if p.suffix.lower() in IMAGE_EXT:
        return llm.ask(question, images=[p], system=system, task="vision")
    text = extract_text(p)
    header = f"File: {p.name}\nQuestion: {question}\n\n"
    if llm.gemini_available() and len(text) < 600_000:
        try:
            return llm.ask(header + "CONTENT:\n" + text, system=system, task="analyze")
        except Exception:
            pass
    # Local model has a small context window: map-reduce over chunks.
    size = 12000
    parts = _chunks(text, size)
    if len(parts) == 1:
        return llm.ask(header + "CONTENT:\n" + text, system=system, task="analyze", local_only=True)
    notes = []
    for i, ch in enumerate(parts[:40], 1):
        notes.append(llm.ask(f"{header}This is part {i}/{len(parts)}. Extract everything relevant to the question "
                             f"as terse notes with page/section refs.\n\n{ch}", system=system, task="analyze_chunk",
                             local_only=True))
    return llm.ask(header + "Notes from each part of the file:\n\n" + "\n\n".join(notes) +
                   "\n\nNow write the final answer.", system=system, task="analyze", local_only=True)
