"""Slide decks (PowerPoint) with generated images and speaker notes. Same spec powers teacher mode."""
import datetime
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .. import config, ui
from ..llm import router
from .docs import slug
from .images import generate_many
from .registry import tool

BG = RGBColor(0x0B, 0x12, 0x20)
PANEL = RGBColor(0x13, 0x1E, 0x33)
ACCENT = RGBColor(0x35, 0xC8, 0xFF)
TEXT = RGBColor(0xF2, 0xF5, 0xFA)
MUTED = RGBColor(0x9A, 0xA8, 0xBF)


def deck_spec(topic, n_slides=10, source_text="", audience="college students", teaching=False):
    """Ask the LLM for a structured deck. `notes` doubles as the narration script in teacher mode."""
    narr = ("notes: the teacher's spoken narration for this slide, 90-150 words, conversational like a great professor: "
            "explain, give a concrete example or analogy, occasionally ask the audience a rhetorical question, and "
            "transition to the next slide" if teaching else
            "notes: speaker notes, 60-100 words, what the presenter should say")
    src = f"\n\nBase the content on this material (cite it where relevant):\n{source_text[:120000]}" if source_text else ""
    spec = router().ask_json(
        f"Create a {n_slides}-slide presentation on: {topic}\nAudience: {audience}.{src}\n\n"
        "Structure: a clear narrative arc (hook → core ideas → examples/evidence → recap). Slides must be concise: "
        "3-5 bullets each, max 12 words per bullet, no paragraphs on slides.\n"
        "Return JSON: {\"title\": str, \"subtitle\": str, \"slides\": [{\"title\": str, \"bullets\": [str], "
        f"\"{narr.split(':')[0]}\": str, \"image_prompt\": str (a vivid description of an illustrative image for this "
        "slide, or \"\" if an image wouldn't help)}], \"quiz\": [{\"question\": str, \"options\": [str], \"answer\": int}] "
        f"(3 multiple-choice questions to check understanding)}}\nRules for {narr}",
        system="You are an expert educator and presentation designer. Be accurate; don't invent statistics.",
        task="deck_teaching" if teaching else "deck", temperature=0.5)
    spec.setdefault("title", topic)
    spec["slides"] = [s for s in spec.get("slides", []) if isinstance(s, dict) and s.get("title")]
    for s in spec["slides"]:
        s["bullets"] = [str(b) for b in (s.get("bullets") or [])][:6]
        s.setdefault("notes", "")
    return spec


def _bg(slide, prs):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _text(slide, x, y, w, h, text, size, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text, p.alignment = text, align
    p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(size), bold, color, "Helvetica Neue"
    return tb


def build_pptx(spec, out_path, images=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height
    images = images or {}

    s = prs.slides.add_slide(blank)
    _bg(s, prs)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.55), Inches(1.4), Emu(50000))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
    _text(s, Inches(0.8), Inches(1.7), Inches(11.5), Inches(1.8), spec["title"], 44, bold=True)
    _text(s, Inches(0.8), Inches(3.8), Inches(11.5), Inches(1.2), spec.get("subtitle", ""), 22, MUTED)
    _text(s, Inches(0.8), Inches(6.6), Inches(11.5), Inches(0.5),
          f"{config.USER_NAME} · {datetime.date.today():%B %Y}", 12, MUTED)
    if spec.get("intro_notes"):
        s.notes_slide.notes_text_frame.text = spec["intro_notes"]

    for i, sd in enumerate(spec["slides"]):
        s = prs.slides.add_slide(blank)
        _bg(s, prs)
        img = images.get(i)
        text_w = Inches(7.0) if img else Inches(11.7)
        _text(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0), sd["title"], 32, bold=True)
        ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(0.9), Emu(40000))
        ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.8), Inches(1.8), text_w, Inches(5.2))
        tf = tb.text_frame
        tf.word_wrap = True
        for j, b in enumerate(sd["bullets"]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = f"▸  {b}"
            p.font.size, p.font.color.rgb, p.font.name = Pt(22), TEXT, "Helvetica Neue"
            p.space_after = Pt(14)
        if img and Path(img).exists():
            s.shapes.add_picture(str(img), Inches(8.1), Inches(1.8), width=Inches(4.6))
        _text(s, Inches(12.2), Inches(6.9), Inches(0.9), Inches(0.4), str(i + 2), 11, MUTED, align=PP_ALIGN.RIGHT)
        s.notes_slide.notes_text_frame.text = sd.get("notes", "")

    if spec.get("references"):
        s = prs.slides.add_slide(blank)
        _bg(s, prs)
        _text(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0), "References", 32, bold=True)
        tb = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.5))
        tf = tb.text_frame
        tf.word_wrap = True
        for j, r in enumerate(spec["references"][:12]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = r
            p.font.size, p.font.color.rgb = Pt(11), MUTED
    prs.save(out_path)
    return str(out_path)


def render_images(spec, folder):
    folder = Path(folder)
    jobs = [(sd.get("image_prompt", ""), folder / f"slide_{i:02d}") for i, sd in enumerate(spec["slides"])]
    ui.status(f"Generating {sum(1 for j in jobs if j[0])} images")
    res = generate_many(jobs)
    return {i: p for i, p in enumerate(res) if p}


@tool("Create a PowerPoint slide deck on any topic (optionally from a file or research report), with images and "
      "speaker notes. Opens it when done.",
      topic="deck topic", slides="number of content slides", source_file="optional file to base it on",
      audience="who it's for", images="generate illustrations")
def make_slides(topic, slides=10, source_file="", audience="college students", images=True):
    from .files import extract_text
    src = extract_text(source_file) if source_file else ""
    ui.say(f"Designing a {slides}-slide deck on {topic}.")
    spec = deck_spec(topic, int(slides), src, audience)
    name = f"{slug(topic)}_{datetime.datetime.now():%Y%m%d_%H%M}"
    folder = config.OUTPUT_DIR / "slides" / name
    folder.mkdir(parents=True, exist_ok=True)
    imgs = render_images(spec, folder) if images else {}
    out = build_pptx(spec, config.OUTPUT_DIR / f"{name}.pptx", imgs)
    (folder / "spec.json").write_text(json.dumps(spec, indent=2))
    import subprocess
    subprocess.run(["open", out])
    return f"Slides saved: {out}"


def slides_from_research(json_path, n_slides=10):
    data = json.loads(Path(json_path).read_text())
    syn = data["synthesis"]
    material = json.dumps({"synthesis": syn, "sources": [
        {"n": i, "title": p["title"], "summary": p.get("summary"), "findings": p.get("key_findings")}
        for i, p in enumerate(data["papers"], 1)]})
    spec = deck_spec(data["topic"], n_slides, "Research brief (cite as [n]):\n" + material)
    spec["references"] = [f"[{i}] {p['citation']}" for i, p in enumerate(data["papers"], 1)]
    folder = Path(json_path).with_suffix("")
    folder.mkdir(exist_ok=True)
    out = build_pptx(spec, Path(json_path).with_suffix(".pptx"), render_images(spec, folder))
    return out
